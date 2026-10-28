import copy
from base64 import b64decode
from io import BytesIO
from itertools import takewhile
import subprocess
import sys
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from api import docs_other as reports


def read_report(encoded):
    return load_workbook(BytesIO(b64decode(encoded, validate=True)))


def service_inputs():
    edges = [{
        "id": 1, "name": "edge-a", "softwareVersion": "5.2.3",
        "modelNumber": "EDGE620", "haState": "UNCONFIGURED",
        "haMode": "NONE", "edgeState": "CONNECTED",
        "licenses": [{"alias": "PREMIUM | 100 Mbps"}],
        "site": {"timezone": "Europe/Prague", "city": "Prague"},
        "wanLinks": {"links": [
            {"internalId": "link-a", "name": "Internet", "mode": "PUBLIC",
             "isp": "carrier", "interfaces": ["GE1", "GE2"]},
            {"internalId": "link-b", "name": "Private", "mode": "PRIVATE",
             "isp": "private carrier", "interfaces": "GE3"},
        ]},
    }]
    performance = [
        {"internalId": "link-a", "bestLatencyMsRx": 101, "bestLatencyMsTx": 50,
         "bestJitterMsRx": 11, "bestJitterMsTx": 5,
         "bestLossPctRx": 1.1, "bestLossPctTx": 0.1,
         "bytesRx": 2048, "bytesTx": 1024},
        {"linkLogicalId": "link-b", "bestLatencyMsRx": 50, "bestLatencyMsTx": 102,
         "bestJitterMsRx": 5, "bestJitterMsTx": 12,
         "bestLossPctRx": 0.1, "bestLossPctTx": 2,
         "bytesRx": 1024, "bytesTx": 4096},
    ]
    return edges, performance


def capacity_inputs():
    paths = []
    for index, utilization in enumerate((0, 5, 20)):
        row = {
            "edge_name": "edge-a", "link_name": f"link-{index}",
            "interface": "GE1", "isp": "carrier",
            "Avg Throughput Rx [Mbps]": 10, "Avg Throughput Tx [Mbps]": 20,
        }
        for direction in ("Rx", "Tx"):
            for bucket, value in (("80-100%", utilization), ("60-80%", 0),
                                  ("30-60%", 0), ("<30%", 100 - utilization)):
                row[f"Util {bucket} {direction} [%]"] = value
                row[f"Util {bucket} {direction} BH [%]"] = value
        paths.append(row)
    aggregated = {
        1: {"edge_id": 1, "edge_name": "edge-a", "link_count": 3,
            "all_throughput_values": list(range(1, 21))},
        2: {"edge_id": 2, "edge_name": "edge-b", "link_count": 0,
            "all_throughput_values": []},
    }
    edges = [
        {"id": 1, "modelNumber": "EDGE620", "licenses": [{"alias": "PREMIUM | 10 Mbps"}]},
        {"id": 2, "modelNumber": "EDGE610", "licenses": []},
    ]
    return paths, aggregated, edges


def sheet_records(sheet):
    rows = list(sheet.values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


class ServiceManagerReportTests(unittest.TestCase):
    def test_legacy_name_is_compatible(self):
        self.assertIs(reports.Service_Manager_Report, reports.service_manager_report)

    def test_parsing_precedence_and_units(self):
        self.assertEqual(reports._service_release_dates("4.2 then 5.2.3"),
                         ("2027-03-16", "LTS release"))
        for value, expected in (
            ("100. Mbps", "100. Mbps"),
            ("1.25\tGbps", "1.25 Gbps"),
            ("version 10.2.3 Mbps", "2.3 Mbps"),
            ("100 mbps", "Unknown"),
            ("unknown", "Unknown"),
        ):
            with self.subTest(value=value):
                self.assertEqual(reports._service_license_bandwidth(value), expected)

    def test_long_nonmatching_numeric_input_is_bounded(self):
        script = (
            "from api.docs_other import _service_release_dates, _service_license_bandwidth\n"
            "text = '9' * 300000\n"
            "assert _service_release_dates(text) == ('', '')\n"
            "assert _service_release_dates(text + '.1') == ('', '')\n"
            "assert _service_license_bandwidth(text) == 'Unknown'\n"
            "assert _service_license_bandwidth(text + '.' + text) == 'Unknown'\n"
        )
        subprocess.run([sys.executable, "-c", script], check=True, timeout=10,
                       capture_output=True, text=True)

    def test_workbook_contract(self):
        edges, performance = service_inputs()
        original = copy.deepcopy((edges, performance))
        workbook = read_report(reports.Service_Manager_Report(edges, performance))
        self.assertEqual((edges, performance), original)
        self.assertEqual(workbook.sheetnames, [
            "Sites", "Circuits", "data_model", "data_sw", "data_license",
            "data_bandwidth", "data_ha", "Analysis", "Summary",
        ])
        for name in workbook.sheetnames[2:7]:
            self.assertEqual(workbook[name].sheet_state, "hidden")
        site = sheet_records(workbook["Sites"])[0]
        self.assertEqual(site["Hostname"], "edge-a")
        self.assertEqual(site["End of General Support"], "2027-03-16")
        self.assertEqual(site["End of Technical Guidance"], "LTS release")
        self.assertEqual(site["License Level"], "PREMIUM")
        self.assertEqual(site["License Bandwidth"], "100 Mbps")
        self.assertIsNone(site["HA Serial Number"])
        circuits = sheet_records(workbook["Circuits"])
        self.assertEqual(circuits[0]["interfaces"], "GE1, GE2")
        self.assertEqual(circuits[1]["interfaces"], "GE3")
        self.assertEqual(circuits[0]["Bytes RX"], "2.00 KB")
        self.assertEqual(circuits[1]["Latency TX (ms)"], 102)
        self.assertEqual(len(workbook["Sites"].tables), 1)
        self.assertEqual(len(workbook["Circuits"].tables), 1)
        self.assertEqual(len(workbook["Summary"]._charts), 5)
        chart = workbook["Summary"]._charts[0]
        self.assertIn("data_model!", chart.ser[0].val.numRef.f)
        analysis = {row[0]: row[1:] for row in workbook["Analysis"].values if row[0]}
        self.assertEqual(analysis["Total Sites"][0], 1)
        self.assertEqual(analysis["Total Circuits"][0], 2)
        self.assertEqual(analysis["High Latency"][:2], ("> 50 ms", 1))
        self.assertEqual(analysis["Very High Latency"][:2], ("> 100 ms", 1))
        self.assertEqual(analysis["Packet Loss"][:2], ("> 0.1%", 1))
        self.assertEqual(workbook["Analysis"]["A1"].value,
                         "Network Infrastructure Analysis Dashboard")
        for heading, expected_overlay, color in (
            ("5. TOP 10 WORST CIRCUITS BY LATENCY (RX) TO VELOCLOUD GW", "Internet", "FF6B6B"),
            ("5B. TOP 10 WORST CIRCUITS BY LATENCY (TX) TO VELOCLOUD GW", "Private", "FF6B6B"),
        ):
            sheet = workbook["Analysis"]
            title_row = next(cell.row for row in sheet for cell in row if cell.value == heading)
            self.assertEqual(sheet.cell(title_row + 2, 2).value, expected_overlay)
            self.assertTrue(sheet.cell(title_row + 2, 2).fill.fgColor.rgb.endswith(color))

    def test_release_and_bandwidth_matching(self):
        cases = [
            ("release 5.2.3 build 9", "PREMIUM 1.5 Gbps", "2027-03-16", "1.5 Gbps"),
            ("version 4.2", "STANDARD 100Mbps", "2023-06-30", "100 Mbps"),
            ("unknown", "unknown", None, "Unknown"),
            ("6.1.2", "ENTERPRISE 2 Gbps", "Not yet announced", "2 Gbps"),
        ]
        for version, license_name, support, bandwidth in cases:
            with self.subTest(version=version):
                edges, performance = service_inputs()
                edges[0]["softwareVersion"] = version
                edges[0]["licenses"][0]["alias"] = license_name
                workbook = read_report(reports.Service_Manager_Report(edges, performance))
                site = sheet_records(workbook["Sites"])[0]
                self.assertEqual(site["End of General Support"], support)
                self.assertEqual(site["License Bandwidth"], bandwidth)


class CapacityReportTests(unittest.TestCase):
    def test_calculations_sheets_and_status(self):
        inputs = capacity_inputs()
        original = copy.deepcopy(inputs)
        workbook = read_report(reports.capacity_planner_report(*inputs))
        self.assertEqual(inputs, original)
        self.assertEqual(workbook.sheetnames, [
            "README", "Link Details", "Edge Summary",
            "Link Utilization Analysis", "License Analysis",
        ])
        summary = sheet_records(workbook["Edge Summary"])
        self.assertEqual(summary[0]["Aggregated Max [Mbps]"], 20)
        self.assertEqual(summary[0]["Aggregated Throughput Avg [Mbps]"], 10.5)
        self.assertEqual(summary[0]["Aggregated Throughput 95th percentile [Mbps]"], 20)
        self.assertEqual(summary[0]["License sizing"], "Undersized")
        self.assertEqual(summary[1]["license_BW"], "N/A")
        self.assertEqual(summary[1]["Aggregated Max [Mbps]"], 0)
        cells = [cell.value for row in workbook["Link Utilization Analysis"] for cell in row]
        for status in ("Critical", "High", "Idle"):
            self.assertIn(status, cells)

    def test_empty_report_has_only_reference_sheet(self):
        workbook = read_report(reports.capacity_planner_report([], {}, []))
        self.assertEqual(workbook.sheetnames, ["README"])

    def test_top_n_boundaries(self):
        for count, expected in ((99, 10), (100, 20), (500, 20), (501, 30)):
            with self.subTest(count=count):
                paths, aggregated, edges = capacity_inputs()
                paths = [dict(paths[1], link_name=f"link-{index}") for index in range(35)]
                edges.extend({"id": index + 10} for index in range(count - len(edges)))
                workbook = read_report(reports.capacity_planner_report(paths, aggregated, edges))
                cells = [cell.value for row in workbook["Link Utilization Analysis"] for cell in row]
                heading = f"TOP {expected} MOST CONGESTED LINKS (Business Hours >80%) & Overall >80%"
                self.assertIn(heading, cells)
                rows = list(workbook["Link Utilization Analysis"].values)
                title_index = next(index for index, row in enumerate(rows) if row[0] == heading)
                ranked_rows = list(takewhile(lambda row: row[0] is not None, rows[title_index + 3:]))
                self.assertEqual(len(ranked_rows), expected)

    def test_status_thresholds(self):
        for high, expected in ((20, "Critical"), (19.9, "High"), (5, "High"),
                               (4.9, "Medium"), (0, "OK")):
            with self.subTest(high=high):
                self.assertEqual(reports._get_link_status({"Util 80-100% Rx [%]": high}), expected)
        for value, expected in ((80, "CRITICAL"), (60, "MEDIUM"), (59.9, "LOW")):
            self.assertEqual(reports._get_risk_level({"utilization_pct": value}), expected)


class PolicyHelperTests(unittest.TestCase):
    def test_edge_and_profile_rows_preserve_order_and_raw_sheet(self):
        device_module = {"name": "deviceSettings", "data": {"segments": []}}
        qos_module = {"name": "QOS", "data": {"segments": [
            {"segment": {"name": "Global", "segmentId": 0}, "rules": [{"name": "allow"}]},
        ]}}
        modules = [device_module, qos_module]
        edges = [
            {"name": "edge-a", "configstack": {"modules": modules}},
            {"name": "ignored-edge", "modules": modules},
        ]
        profiles = [{"name": "profile-a", "modules": modules}]
        with patch.object(reports, "parse_business_rule",
                          return_value={"name": "allow", "dst": "prefix: any"}) as parse:
            encoded = reports.danone_bp_extract_generate(
                edges, {"applications": []}, {"applications": []}, [], profiles)
        workbook = read_report(encoded)
        raw = sheet_records(workbook["BPExtract"])
        policy = sheet_records(workbook["Policy Table"])
        self.assertEqual(parse.call_count, 2)
        self.assertEqual([row["Source Type"] for row in raw], ["edge", "profile"])
        self.assertEqual([row["Source Name"] for row in raw], ["edge-a", "profile-a"])
        self.assertEqual([row["Destination"] for row in policy], ["Internet", "Internet"])
        self.assertEqual(workbook["Policy Table"].tables["PolicyRulesView"].ref, "A1:L3")

    def test_reference_order_and_identity(self):
        first = {"logicalId": "one", "name": "first"}
        duplicate = {"logicalId": "one", "name": "duplicate"}
        second = {"name": "second"}
        gateway = {"logicalId": "gateway"}
        device = {"refs": {"deviceSettings:css:provider": [first, second],
                           "deviceSettings:dataCenter": gateway}}
        qos = {"refs": {"deviceSettings:css:provider": [duplicate],
                        "deviceSettings:vpn:dataCenter": [gateway]}}
        self.assertEqual(reports._collect_css_and_nsdgw(device, qos),
                         ([first, second], [gateway]))

    def test_recursive_lookup_preserves_precedence(self):
        data = {"logicalId": "dc:one", "name": "first",
                "nested": [{"id": "one", "name": "second"}]}
        self.assertEqual(reports._find_name_for_logical_id(data, "one"), "first")
        self.assertEqual(reports._find_name_for_logical_id([data], "missing"), "")
        self.assertEqual(reports._find_name_for_logical_id(
            {"nsdLogicalId": "one", "nsd": "not previously recognized"}, "one"), "")

    def test_cidr_forms(self):
        cases = [
            ("192.0.2.129", "255.255.255.0", "192.0.2.0/24"),
            ("192.0.2.1", "", "192.0.2.1/32"),
            ("2001:db8::1", "64", "2001:db8::/64"),
            ("2001:db8::1", "ffff:ffff:ffff:ffff::", "2001:db8::/64"),
            ("2001:db8::1", "ffff:0:ffff::", ""),
            ("invalid", "24", ""),
            ("192.0.2.129/24", None, ""),
        ]
        for address, mask, expected in cases:
            with self.subTest(address=address, mask=mask):
                self.assertEqual(reports._to_cidr(address, mask), expected)

    def test_policy_workbook_structure(self):
        encoded = reports.danone_bp_extract_generate(
            [], {"applications": []}, {"applications": []}, [], [])
        workbook = read_report(encoded)
        self.assertEqual(workbook.sheetnames, ["Policy Table", "BPExtract"])
        self.assertEqual(workbook.active.title, "Policy Table")
        self.assertEqual(workbook["BPExtract"]["A1"].value, "Source Type")
        policy = workbook["Policy Table"]
        self.assertEqual(policy["A1"].value, "Source Name")
        self.assertEqual(policy["L1"].value, "DSCP Outer")
        self.assertEqual(policy.tables["PolicyRulesView"].ref, "A1:L2")
        self.assertFalse(policy.sheet_view.showGridLines)
        self.assertTrue(policy.column_dimensions["M"].hidden)
        self.assertEqual(len(policy.conditional_formatting), 1)
        self.assertTrue(policy["A1"].font.bold)
        self.assertTrue(policy["A1"].fill.fgColor.rgb.endswith("17365D"))


class EventReportTests(unittest.TestCase):
    def test_sorting_and_downtime(self):
        events = [
            {"id": 2, "edgeName": "edge-a", "event": "LINK_ALIVE",
             "eventTime": "2026-01-01T10:01:00.000Z",
             "created": "2026-01-01T10:01:01.000Z", "message": "GE1 up",
             "detail": '{"internalId":"one"}'},
            {"id": 1, "edgeName": "edge-a", "event": "LINK_DEAD",
             "eventTime": "2026-01-01T10:00:00.000Z",
             "created": "2026-01-01T10:00:01.000Z", "message": "GE1 down",
             "detail": '{"internalId":"one"}'},
        ]
        workbook = read_report(reports.event_logs_parse(
            events, [{"link": {"internalId": "one", "displayName": "Internet"}}]))
        self.assertEqual(workbook.sheetnames, ["Time Differences", "Downtime Summary"])
        record = sheet_records(workbook["Downtime Summary"])[0]
        self.assertEqual(record["downtime[sec]"], 60)
        self.assertEqual(record["occurrences"], 1)
        self.assertEqual(record["Overlay name"], "Internet")
