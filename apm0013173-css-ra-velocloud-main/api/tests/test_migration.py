"""Differential tests against original source; never contact an orchestrator."""
import ast
import copy
import datetime
import importlib.util
import json
import sys
import types
import zipfile
from base64 import b64decode
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

import django
django.setup()

import pytest
from django.test import Client, override_settings
from django.urls import path
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api import serializers, views
from api.urls import urlpatterns
from velocloud_app.asgi import application


ROOT = Path(__file__).resolve().parents[2]
BASELINE = ROOT / ".migration-baseline"


def original_source(path):
    with zipfile.ZipFile(ROOT / "api/tests/fixtures/django_baseline.zip") as archive:
        return archive.read(path).decode("utf-8").replace("\r\n", "\n")


@pytest.fixture(scope="module")
def original():
    module = types.ModuleType("_original_serializers")
    exec(compile(original_source("api/serializers.py"), "original/api/serializers.py", "exec"), module.__dict__)
    sys.modules[module.__name__] = module
    original_views = types.ModuleType("_original_views")
    source = original_source("api/views.py").replace(
        "from api import serializers", "import _original_serializers as serializers")
    exec(compile(source, "original/api/views.py", "exec"), original_views.__dict__)
    sys.modules[original_views.__name__] = original_views
    api_urls = types.ModuleType("_original_api_urls")
    source = original_source("api/urls.py").replace("from api import views", "import _original_views as views")
    exec(compile(source, "original/api/urls.py", "exec"), api_urls.__dict__)
    from django.urls import include
    from django.views.generic.base import RedirectView
    from django.views.static import serve
    from functools import partial
    urls = types.ModuleType("_original_urls")
    urls.urlpatterns = [path("api/", include(api_urls)),
                        path("", RedirectView.as_view(pattern_name="schema-swagger-ui", permanent=True)),
                        path("static/<path:path>", partial(serve, document_root=str(ROOT / "velocloud_app/static")))]
    return types.SimpleNamespace(serializers=module, views=original_views, urls=urls)


@pytest.fixture(scope="module")
def fastapi_client():
    with TestClient(application, base_url="http://localhost", raise_server_exceptions=False) as client:
        yield client


def django_call(original, method, route, body=b"", content_type="application/json", headers=None):
    with override_settings(ROOT_URLCONF=original.urls):
        response = Client(raise_request_exception=False).generic(
            method, route, body, content_type=content_type, HTTP_HOST="localhost", **{
                "HTTP_" + key.upper().replace("-", "_"): value for key, value in (headers or {}).items()})
        if response.streaming:
            try:
                return types.SimpleNamespace(status_code=response.status_code, headers=response.headers,
                    content=b"".join(response.streaming_content), cookies=response.cookies)
            finally:
                response.close()
        return response


def assert_http_equal(before, after):
    assert after.status_code == before.status_code
    if before.headers.get("content-type", "").startswith("text/html") and b"csrfmiddlewaretoken" in before.content:
        assert normalize_html(after.content) == normalize_html(before.content)
    else:
        assert after.content == before.content
    for header in ("content-type", "allow", "vary", "www-authenticate", "x-frame-options",
                   "x-content-type-options", "referrer-policy", "cross-origin-opener-policy"):
        assert after.headers.get(header) == before.headers.get(header), header


def assert_cookies_equal(before, after):
    assert sorted(after.headers.get_list("set-cookie")) == sorted(
        cookie.OutputString() for cookie in before.cookies.values())


def test_business_source_is_unchanged():
    for file in ("utils.py", "docs_mds.py", "docs_other.py", "doc_utils_helpers.py"):
        assert (ROOT / "api" / file).read_text(encoding="utf-8") == original_source("api/" + file)
    old = ast.parse(original_source("api/views.py"))
    new = ast.parse((ROOT / "api/views.py").read_text(encoding="utf-8"))
    def functions(tree):
        return [ast.dump(node, include_attributes=False) for node in ast.walk(tree)
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))]
    assert functions(old) == functions(new)


@pytest.mark.parametrize("name", ["RAGenericInput", "MdsSerializer", "SingleEdgeProvisioning", "HWStatsSerializer"])
def test_serializer_differential(original, name):
    seeds = {
        "RAGenericInput": {"params": [{"foo": [1, True, {"x": 1}]}], "devices": [{"hostname": 123}], "output": "  json  "},
        "MdsSerializer": {"vco_url": "example.velocloud.net", "vco_username": "u", "vco_password": "p"},
        "SingleEdgeProvisioning": {"routedInterfaces": [], "links": [], "description": "edge", "site": {}, "bgp": {}},
        "HWStatsSerializer": {"startDate": "2026-10-01", "endDate": "2026-10-02", "customerName": "example"},
    }
    seed = seeds[name]
    cases = [None, {}, [], "value", 42, seed, dict(seed, unknown="ignored")]
    for field in seed:
        missing = dict(seed)
        del missing[field]
        cases.append(missing)
        cases.extend(dict(seed, **{field: value}) for value in (None, "", " ", True, 0, 1.5, [], {}, [None], [1]))
    if name == "RAGenericInput":
        cases += [dict(seed, output=value) for value in (None, "", "\x00", True, {})]
    if name == "HWStatsSerializer":
        cases += [dict(seed, endDate="2026-09-30"), dict(seed, startDate="2026-2-1"),
                  dict(seed, startDate="2026-02-30"), dict(seed, output="JSON")]
    for data in cases:
        expected = getattr(original.serializers, name)(data=copy.deepcopy(data))
        actual = getattr(serializers, name)(data=copy.deepcopy(data))
        assert actual.is_valid() == expected.is_valid(), (name, data, actual.errors, expected.errors)
        assert actual.errors == expected.errors, (name, data)
        assert actual.validated_data == expected.validated_data, (name, data)
        if expected.is_valid():
            assert actual.data == expected.data, (name, data)


def test_many_serializer_differential(original):
    valid = {"routedInterfaces": [None], "links": [1], "description": "edge", "site": {}, "bgp": {}}
    for data in ([], [valid], [valid, {}], None, {}, [None], [dict(valid, old_edge=" ")]):
        expected = original.serializers.SingleEdgeProvisioning(data=data, many=True)
        actual = serializers.SingleEdgeProvisioning(data=data, many=True)
        assert actual.is_valid() == expected.is_valid()
        assert actual.errors == expected.errors
        assert actual.validated_data == expected.validated_data


def test_all_registered_route_validation(original, fastapi_client):
    for route, cls in urlpatterns:
        before = django_call(original, "POST", "/api/" + route, b"{}")
        after = fastapi_client.post("/api/" + route, content=b"{}", headers={"content-type": "application/json"})
        assert_http_equal(before, after)
        before_options = django_call(original, "OPTIONS", "/api/" + route)
        assert_http_equal(before_options, fastapi_client.options("/api/" + route))


def normalize_html(value):
    import re
    return re.sub(r"[A-Za-z0-9]{64}", "CSRF_TOKEN", value.decode())


def test_browser_and_swagger_contracts(original, fastapi_client):
    import yaml
    for route in ("/api/swagger.json", "/api/swagger.yaml", "/api/swagger/", "/api/redoc/"):
        before = django_call(original, "GET", route)
        after = fastapi_client.get(route)
        assert before.status_code == after.status_code == 200
        assert before.headers["content-type"] == after.headers["content-type"]
        if route.endswith("json"):
            assert json.loads(before.content) == after.json()
        elif route.endswith("yaml"):
            assert yaml.safe_load(before.content) == yaml.safe_load(after.content)
        else:
            assert normalize_html(before.content) == normalize_html(after.content)
    for route in ("/api/report/billing", "/api/mds/", "/api/report/hwstats"):
        before = django_call(original, "GET", route, headers={"Accept": "text/html"})
        after = fastapi_client.get(route, headers={"Accept": "text/html"})
        assert before.status_code == after.status_code == 405
        assert before.headers["content-type"] == after.headers["content-type"]
        assert normalize_html(before.content) == normalize_html(after.content)
        assert "csrftoken" in after.cookies
    for asset in ("rest_framework/js/jquery-3.7.1.min.js", "drf-yasg/swagger-ui-init.js", "drf-yasg/redoc/redoc.min.js"):
        before = django_call(original, "GET", "/static/" + asset)
        after = fastapi_client.get("/static/" + asset)
        assert_http_equal(before, after)
        for header in ("content-disposition", "last-modified", "content-length"):
            assert before.headers.get(header) == after.headers.get(header)
    for route in ("/api/swagger/?format=openapi", "/api/swagger/?format=.json", "/api/swagger/?format=json",
                  "/api/swagger.json?format=.yaml", "/api/swagger/?format=redoc"):
        before = django_call(original, "GET", route)
        after = fastapi_client.get(route)
        assert_http_equal(before, after)
    for method in ("POST", "OPTIONS", "HEAD"):
        assert_http_equal(django_call(original, method, "/api/swagger.json"),
                          fastapi_client.request(method, "/api/swagger.json"))


def test_root_csrf_and_api_exemption(original, fastapi_client):
    secret = "a" * 32
    cases = [{}, {"Cookie": "csrftoken=" + secret},
             {"Cookie": "csrftoken=" + secret, "X-CSRFToken": secret},
             {"Cookie": "csrftoken=" + secret, "X-CSRFToken": "b" * 32},
             {"Cookie": "csrftoken=" + secret, "X-CSRFToken": secret, "Origin": "https://untrusted.invalid"},
             {"Cookie": "csrftoken=" + secret, "X-CSRFToken": "bad"}]
    for debug in (True, False):
        from velocloud_app import settings as runtime_settings
        before_debug = runtime_settings.DEBUG
        runtime_settings.DEBUG = debug
        try:
            with override_settings(ROOT_URLCONF=original.urls, DEBUG=debug):
                for headers in cases:
                    fastapi_client.cookies.clear()
                    before = Client(enforce_csrf_checks=True).generic("POST", "/", b"", content_type="application/json", HTTP_HOST="localhost",
                        **{"HTTP_" + key.upper().replace("-", "_"): value for key, value in headers.items()})
                    after = fastapi_client.post("/", headers=headers, follow_redirects=False)
                    assert_http_equal(before, after)
        finally:
            runtime_settings.DEBUG = before_debug
    response = fastapi_client.post("/api/report/billing", json={}, headers={"Origin": "https://untrusted.invalid"})
    assert response.status_code == 400


@pytest.mark.parametrize("method,body,content_type,headers", [
    ("GET", b"", "application/json", {}), ("HEAD", b"", "application/json", {}),
    ("PUT", b"{", "application/json", {}), ("OPTIONS", b"", "application/json", {}),
    ("POST", b"{", "application/json", {}), ("POST", b"", "application/json", {}),
    ("POST", b"null", "application/json", {}), ("POST", b"[]", "application/json", {}),
    ("POST", b"{}", "text/plain", {}), ("POST", b"{}", "application/json", {"Accept": "application/xml"}),
    ("POST", b"{", "application/json", {"Authorization": "Token"}),
    ("POST", b"{", "application/json", {"Authorization": "Token one two"}),
    ("POST", b"{}", "application/json", {"Authorization": "Basic dTpw"}),
    ("GET", b"", "application/json", {"Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"}),
    ("POST", b"{}", "application/json", {"Accept": "application/json;indent=4,text/html"}),
    ("POST", b"{}", "application/x-www-form-urlencoded", {}),
])
def test_http_boundary(original, fastapi_client, method, body, content_type, headers):
    before = django_call(original, method, "/api/report/billing", body, content_type, headers)
    after = fastapi_client.request(method, "/api/report/billing", content=body,
                                   headers={"Content-Type": content_type, **headers})
    assert_http_equal(before, after)


def test_form_multipart_and_encoding_parity(original, fastapi_client):
    from django.test.client import encode_multipart, BOUNDARY
    from django.core.files.uploadedfile import SimpleUploadedFile
    cases = [("/api/report/billing", b"params=[]&devices=[]", "application/x-www-form-urlencoded"),
             ("/api/report/billing", b"params[0]foo=1", "application/x-www-form-urlencoded"),
             ("/api/mds/", b"vco_url=example.velocloud.net&vco_username=&vco_password=p&emails=a&emails=b", "application/x-www-form-urlencoded"),
             ("/api/report/billing", b"output=\xff", "application/x-www-form-urlencoded"),
             ("/api/report/billing", b"{}", ""),
             ("/api/report/billing", b'{"x":"\xff"}', "application/json"),
             ("/api/report/billing", b"{}", "application/json; charset=unknown-encoding"),
             ("/api/report/billing", b"invalid", "multipart/form-data"),
             ("/api/report/billing", b"invalid", "multipart/form-data; boundary=foo")]
    for values in ({"params": "[]", "devices": "[]"}, {"output": "x" * 1_100_000},
                   {"params": SimpleUploadedFile("test.txt", b"value")}):
        cases.append(("/api/report/billing", encode_multipart(BOUNDARY, values), "multipart/form-data; boundary=" + BOUNDARY))
    for route, body, content_type in cases:
        before = django_call(original, "POST", route, body, content_type)
        after = fastapi_client.post(route, content=body, headers={"content-type": content_type})
        assert_http_equal(before, after)


def test_browser_validation_errors(original, fastapi_client):
    data = {"startDate": "bad", "endDate": "2026-10-02", "customerName": ""}
    body = json.dumps(data).encode()
    before = django_call(original, "POST", "/api/report/hwstats", body, headers={"Accept": "text/html"})
    after = fastapi_client.post("/api/report/hwstats", content=body,
                                headers={"content-type": "application/json", "Accept": "text/html"})
    assert before.status_code == after.status_code == 400
    assert normalize_html(before.content) == normalize_html(after.content)


def upstream_session(responses):
    session = Mock()
    session.post.side_effect = [Mock(status_code=200, json=Mock(return_value=value)) for value in responses]
    return session


HEADERS = {"X-username": "test-user", "X-password": "test-only", "X-url": "example.velocloud.net", "X-tenant-id": "1"}


def test_mds_real_argument_mismatch(original, fastapi_client):
    body = json.dumps({"vco_url": "example.velocloud.net", "vco_username": "u", "vco_password": "p"}).encode()
    results, calls = [], []
    for framework in ("django", "fastapi"):
        session = upstream_session([{}, {"id": 1}, [], {"applications": []}, {"applications": []}])
        with patch("api.utils.req.Session", return_value=session):
            result = (django_call(original, "POST", "/api/mds/", body) if framework == "django" else
                      fastapi_client.post("/api/mds/", content=body, headers={"content-type": "application/json"}))
        results.append(result)
        calls.append(session.post.call_args_list)
    assert_http_equal(*results)
    assert results[1].status_code == 200
    assert b"missing 2 required positional arguments: 'profiles' and 'segs'" in results[1].content
    assert calls[0] == calls[1]
    # The original business function is never replaced; the TypeError happens
    # after exactly the original login/data fetch sequence.
    assert views.mds_generate is original.views.mds_generate


def test_mds_report_http_workbook_parity(original, fastapi_client):
    class FrozenDatetime(datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 10, 28, 12, tzinfo=tz)
    results, calls = [], []
    with patch("api.utils.datetime.datetime", FrozenDatetime):
        for framework in ("django", "fastapi"):
            session = upstream_session([{}, [], {"applications": [], "applicationClasses": {}},
                                        {"applications": []}, [], []])
            with patch("api.utils.req.Session", return_value=session):
                body = b'{"params":[],"devices":[]}'
                response = (django_call(original, "POST", "/api/report/mds", body, headers=HEADERS)
                    if framework == "django" else fastapi_client.post("/api/report/mds", content=body,
                        headers={"content-type": "application/json", **HEADERS}))
            results.append(json.loads(response.content))
            calls.append(session.post.call_args_list)
            assert response.status_code == 200
    assert calls[0] == calls[1] and len(calls[0]) == 6
    old_encoded = results[0]["data"][0].pop("value")
    new_encoded = results[1]["data"][0].pop("value")
    assert results[0] == results[1]
    before = load_workbook(BytesIO(b64decode(old_encoded)))
    after = load_workbook(BytesIO(b64decode(new_encoded)))
    assert before.sheetnames == after.sheetnames and len(after.sheetnames) == 8
    for old_sheet, new_sheet in zip(before, after):
        assert list(old_sheet.values) == list(new_sheet.values)
        assert [cell.style_id for row in old_sheet for cell in row] == [cell.style_id for row in new_sheet for cell in row]
        assert list(old_sheet.merged_cells.ranges) == list(new_sheet.merged_cells.ranges)


def test_hwstats_real_business_and_workbook(original, fastapi_client):
    for output in ("json", "xlsx"):
        body = json.dumps({"data": {"startDate": "2026-10-01", "endDate": "2026-10-02",
                                  "customerName": "example", "output": output}}).encode()
        results, queries = [], []
        for framework in ("django", "fastapi"):
            cursor = Mock(column_names=("hostname", "datetime", "cpu", "memory"))
            cursor.fetchall.return_value = [("edge", datetime.datetime(2026, 10, 1), 25, 71)]
            connection = Mock()
            connection.cursor.return_value = cursor
            with patch.dict("os.environ", VELO_DB_HOST="test-db", VELO_DB_USER="test", VELO_DB_PASS="test-only"), \
                    patch("api.views.mysql.connector.connect", return_value=connection):
                result = (django_call(original, "POST", "/api/report/hwstats", body) if framework == "django" else
                          fastapi_client.post("/api/report/hwstats", content=body, headers={"content-type": "application/json"}))
            results.append(result)
            queries.append(cursor.execute.call_args)
            cursor.close.assert_called_once()
            connection.close.assert_called_once()
        assert queries[0] == queries[1]
        before, after = [json.loads(result.content) for result in results]
        if output == "json":
            assert_http_equal(*results)
        else:
            old_value = before["data"][0].pop("value")
            new_value = after["data"][0].pop("value")
            assert before == after
            old_book = load_workbook(BytesIO(b64decode(old_value)))
            new_book = load_workbook(BytesIO(b64decode(new_value)))
            assert old_book.sheetnames == new_book.sheetnames
            for old_sheet, new_sheet in zip(old_book, new_book):
                assert list(old_sheet.values) == list(new_sheet.values)
                assert [cell.style_id for row in old_sheet for cell in row] == [cell.style_id for row in new_sheet for cell in row]


@pytest.fixture(scope="module")
def isolated_auth_database():
    from django.conf import settings
    from django.core.management import call_command
    from django.db import connections
    from velocloud_app import settings as runtime_settings
    database = ROOT / "migration_docs/test-results/auth-test.sqlite3"
    database.parent.mkdir(parents=True, exist_ok=True)
    if database.exists():
        database.unlink()
    original_name = settings.DATABASES["default"]["NAME"]
    runtime_name = runtime_settings.DATABASES["default"]["NAME"]
    connections.close_all()
    settings.DATABASES["default"]["NAME"] = str(database)
    runtime_settings.DATABASES["default"]["NAME"] = str(database)
    connections["default"].settings_dict["NAME"] = str(database)
    call_command("migrate", interactive=False, verbosity=0, skip_checks=True)
    yield database
    connections.close_all()
    settings.DATABASES["default"]["NAME"] = original_name
    runtime_settings.DATABASES["default"]["NAME"] = runtime_name
    connections["default"].settings_dict["NAME"] = original_name
    database.unlink()


def test_token_and_session_state(original, fastapi_client, isolated_auth_database):
    from django.contrib.auth import get_user_model
    from django.contrib.sessions.backends.db import SessionStore
    from rest_framework.authtoken.models import Token
    user = get_user_model().objects.create_user(username="migration-user", password="test-only")
    token = Token.objects.create(user=user)
    inactive = get_user_model().objects.create_user(username="inactive-user", is_active=False)
    inactive_token = Token.objects.create(user=inactive)
    session = SessionStore()
    session["_auth_user_id"] = str(user.pk)
    session["_auth_user_backend"] = "django.contrib.auth.backends.ModelBackend"
    session["_auth_user_hash"] = user.get_session_auth_hash()
    session.save()
    remote_session = SessionStore()
    remote_session.update(dict(session.items(), _auth_user_backend="django.contrib.auth.backends.RemoteUserBackend"))
    remote_session.save()
    for headers in ({"Authorization": "Token " + token.key}, {"Authorization": "Token unknown"},
                    {"Authorization": "Token " + inactive_token.key},
                    {"X-Authorization": "Token unknown"},
                    {"Cookie": "sessionid=" + session.session_key}, {"Cookie": "sessionid=invalid"}):
        before = django_call(original, "POST", "/api/report/billing", b"{}", headers=headers)
        after = fastapi_client.post("/api/report/billing", json={}, headers=headers)
        assert_http_equal(before, after)
        assert_cookies_equal(before, after)
    # Both stacks remove a remote-backend session when REMOTE_USER disappears.
    # Recreate it between requests so each framework sees the same valid state.
    before = django_call(original, "POST", "/api/report/billing", b"{}",
                         headers={"Cookie": "sessionid=" + remote_session.session_key})
    remote_session.save(must_create=True)
    after = fastapi_client.post("/api/report/billing", json={},
                                headers={"Cookie": "sessionid=" + remote_session.session_key})
    assert_http_equal(before, after)
    assert_cookies_equal(before, after)
    for route in ("/api/swagger/", "/api/report/billing"):
        headers = {"Authorization": "Token " + token.key, "Accept": "text/html"}
        before = django_call(original, "GET", route, headers=headers)
        after = fastapi_client.get(route, headers=headers)
        assert before.status_code == after.status_code
        assert normalize_html(before.content) == normalize_html(after.content)
    before = django_call(original, "POST", "/api/report/billing", b"{", headers={"Authorization": "Token first,Token second"})
    after = fastapi_client.post("/api/report/billing", content=b"{", headers=[("Content-Type", "application/json"),
                                    ("Authorization", "Token first"), ("Authorization", "Token second")])
    assert_http_equal(before, after)


def test_rma_upstream_payload_and_order(original, fastapi_client):
    inventory = [{"id": 2, "name": "source", "configuration": {
        "enterprise": {"id": 10}, "operator": {"id": 4, "name": "6.0"}}, "licenses": [{"licenseId": 5}]}]
    body = json.dumps({"params": [{"RmaDeviceModel": "7x0"}], "devices": [{"hostname": "source"}]}).encode()
    outcomes, sequences = [], []
    for framework in ("django", "fastapi"):
        # A controlled creation rejection must stop OS/config mutations. All
        # business helpers and payload construction execute in both frameworks.
        session = upstream_session([{}, copy.deepcopy(inventory), {"error": {"message": "creation rejected"}}])
        with patch("api.utils.req.Session", return_value=session):
            result = (django_call(original, "POST", "/api/edge/EdgeRMA", body, headers=HEADERS) if framework == "django" else
                      fastapi_client.post("/api/edge/EdgeRMA", content=body, headers={"content-type": "application/json", **HEADERS}))
        outcomes.append(result)
        sequences.append(session.post.call_args_list)
    assert_http_equal(*outcomes)
    assert sequences[0] == sequences[1]
    assert len(sequences[1]) == 3
    assert sequences[1][2].kwargs["json"]["name"] == "source_RMA"
    assert sequences[1][2].kwargs["json"]["modelNumber"] == "edge7X0"


def test_ingress_prefix_and_direct_requests(fastapi_client, monkeypatch):
    import velocloud_app.asgi as asgi
    monkeypatch.setattr(asgi, "URL_PREFIX", "/svc/velocloud")
    with override_settings(FORCE_SCRIPT_NAME="/svc/velocloud"):
        direct = fastapi_client.post("/api/report/billing", json={})
        prefixed = fastapi_client.post("/svc/velocloud/api/report/billing", json={})
        assert direct.status_code == prefixed.status_code == 400
        assert direct.content == prefixed.content
        for root in ("/", "/svc/velocloud", "/svc/velocloud/"):
            response = fastapi_client.get(root, follow_redirects=False)
            assert response.status_code == 301
            assert response.headers["location"] == "/svc/velocloud/api/swagger/"
        assert fastapi_client.get("/svc/velocloudx/api/report/billing").status_code == 404


def test_real_original_wsgi_and_native_asgi_prefix(original, fastapi_client, monkeypatch):
    import httpx
    from django.test import RequestFactory
    from django.urls import get_script_prefix, set_script_prefix
    from velocloud_app import settings
    import velocloud_app.asgi as asgi
    prefix = "/svc/velocloud"
    monkeypatch.setattr(settings, "FORCE_SCRIPT_NAME", prefix)
    monkeypatch.setattr(asgi, "URL_PREFIX", prefix)
    monkeypatch.setenv("SCRIPT_NAME", prefix)
    old_prefix = get_script_prefix()
    try:
        with override_settings(ROOT_URLCONF=original.urls, FORCE_SCRIPT_NAME=prefix, STATIC_URL=prefix + "/static/"):
            module = types.ModuleType("_original_wsgi")
            exec(compile(original_source("velocloud_app/wsgi.py"), "original/velocloud_app/wsgi.py", "exec"), module.__dict__)
            for uri in ("/api/swagger.json", prefix + "/api/swagger.json", "/api/swagger/", prefix + "/api/swagger/"):
                request = RequestFactory().get(uri, HTTP_HOST="localhost")
                captured = {}
                def start_response(status, headers):
                    captured.update(status_code=int(status.split()[0]), headers=httpx.Headers(headers))
                iterable = module.application(request.META, start_response)
                try:
                    before = types.SimpleNamespace(content=b"".join(iterable), **captured)
                finally:
                    iterable.close()
                after = fastapi_client.get(uri)
                assert_http_equal(before, after)
    finally:
        set_script_prefix(old_prefix)


def test_concurrent_request_isolation(fastapi_client):
    from concurrent.futures import ThreadPoolExecutor
    def send(index):
        body = {"params": []} if index % 2 else {"devices": []}
        response = fastapi_client.post("/api/report/billing", json=body)
        return index, response
    with ThreadPoolExecutor(max_workers=8) as pool:
        for index, response in pool.map(send, range(24)):
            assert response.status_code == 400
            assert response.json() == {"devices" if index % 2 else "params": ["This field is required."]}
    import threading
    import time
    counts = {"active": 0, "peak": 0}
    lock = threading.Lock()
    class Session:
        def post(self, url, **kwargs):
            with lock:
                counts["active"] += 1 if url.endswith("enterpriseLogin") else -1
                counts["peak"] = max(counts["peak"], counts["active"])
            time.sleep(0.005)
            return Mock(status_code=200, json=lambda: {"network": "ok"})
    with patch("api.utils.req.Session", side_effect=Session), ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(lambda _: fastapi_client.post("/api/report/getNetworkOverviewInfo",
                        json={"params": [], "devices": []}, headers=HEADERS), range(8)))
    assert all(result.status_code == 200 for result in results)
    assert counts == {"active": 0, "peak": 1}


def test_deployment_entrypoints():
    for file in ("Dockerfile", "gunicorn.py", "chart/templates/deployment.yaml"):
        source = (ROOT / file).read_text()
        assert "velocloud_app.asgi:application" in source
        assert "uvicorn_worker.UvicornWorker" in source
        assert "velocloud_app.wsgi:application" not in source
    assert "manage.py collectstatic --noinput" in (ROOT / "run.sh").read_text()
    assert "exec gunicorn" in (ROOT / "run.sh").read_text()


def test_signed_records_and_password_hash_compatibility(isolated_auth_database):
    from velocloud_app import auth_state
    from django.contrib.auth import get_user_model
    from django.contrib.sessions.backends.db import SessionStore
    values = [{}, {"message": "accent: café", "list": [1, True, None]}, {"repeat": "value" * 500}]
    with patch("time.time", return_value=1_700_000_000):
        for value in values:
            encoded = SessionStore().encode(value)
            assert auth_state.decode_session(encoded) == value
            assert auth_state.encode_session(value) == encoded
            assert SessionStore().decode(auth_state.encode_session(value)) == value
            assert auth_state.decode_session(encoded[:-1] + ("A" if encoded[-1] != "A" else "B")) == {}
    user = auth_state.create_user("native-created", email="native@EXAMPLE.COM", password="test-only")
    reference = get_user_model().objects.get(pk=user.id)
    assert reference.email == "native@example.com"
    assert reference.check_password("test-only")
    assert not reference.check_password("wrong-password")
    assert reference.get_session_auth_hash() == user.get_session_auth_hash()


def test_original_schema_bootstrap_and_no_rebuild(isolated_auth_database):
    import sqlite3
    from contextlib import closing
    from velocloud_app import auth_state, settings
    fresh = ROOT / "migration_docs/test-results/bootstrap-test.sqlite3"
    if fresh.exists():
        fresh.unlink()
    with closing(sqlite3.connect(isolated_auth_database)) as original_database:
        expected = original_database.execute("SELECT name, sql FROM sqlite_master WHERE sql IS NOT NULL ORDER BY name").fetchall()
    previous = settings.DATABASES["default"]["NAME"]
    settings.DATABASES["default"]["NAME"] = str(fresh)
    try:
        assert auth_state.initialize_database()
        assert not auth_state.initialize_database()
        with closing(sqlite3.connect(fresh)) as database:
            actual = database.execute("SELECT name, sql FROM sqlite_master WHERE sql IS NOT NULL ORDER BY name").fetchall()
            assert actual == expected
            assert database.execute("PRAGMA foreign_key_check").fetchall() == []
        settings.DATABASES["default"]["NAME"] = str(isolated_auth_database)
        assert not auth_state.initialize_database()
        with closing(sqlite3.connect(isolated_auth_database)) as database:
            assert database.execute("SELECT name, sql FROM sqlite_master WHERE sql IS NOT NULL ORDER BY name").fetchall() == expected
    finally:
        settings.DATABASES["default"]["NAME"] = previous
        fresh.unlink()


def test_concurrent_logout_does_not_restore_session(isolated_auth_database):
    from velocloud_app import auth_state
    from django.contrib.sessions.backends.db import SessionStore
    from django.contrib.sessions.backends.base import UpdateError
    original = SessionStore()
    original["value"] = "keep"
    original.save()
    concurrent = SessionStore(original.session_key)
    assert concurrent["value"] == "keep"
    original.delete()
    with pytest.raises(UpdateError):
        concurrent.save()
    native = auth_state.Session(None)
    native.data = {"value": "keep"}
    native.save()
    concurrent_native = auth_state.Session(native.session_key)
    native.flush()
    with pytest.raises(auth_state.SessionInterrupted):
        concurrent_native.save()


def test_server_exception_and_rendering_failure_parity(original, fastapi_client, monkeypatch):
    from velocloud_app import settings
    monkeypatch.setattr(settings, "DEBUG", False)
    monkeypatch.delenv("API_BASE_URL", raising=False)
    fastapi_client.cookies.clear()
    body = b'{"params":[],"devices":[]}'
    with override_settings(DEBUG=False):
        before = django_call(original, "POST", "/api/report/billing", body)
        after = fastapi_client.post("/api/report/billing", content=body, headers={"content-type": "application/json"})
        assert_http_equal(before, after)
        responses = []
        for framework in ("django", "fastapi"):
            session = upstream_session([{}, {"value": float("nan")}])
            with patch("api.utils.req.Session", return_value=session):
                responses.append(django_call(original, "POST", "/api/report/getNetworkOverviewInfo", body, headers=HEADERS)
                    if framework == "django" else fastapi_client.post("/api/report/getNetworkOverviewInfo", content=body,
                        headers={"content-type": "application/json", **HEADERS}))
        assert_http_equal(*responses)
        assert responses[1].status_code == 500
